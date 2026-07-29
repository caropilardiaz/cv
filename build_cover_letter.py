import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = docx.shared.RGBColor(0x33, 0x33, 0x33)

# --- HEADER ---
title_p = doc.add_paragraph()
title_p.paragraph_format.space_after = Pt(0)
r = title_p.add_run("Carolina Del Pilar Díaz")
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = docx.shared.RGBColor(0x00, 0x33, 0x66)

contact_p = doc.add_paragraph()
contact_p.paragraph_format.space_after = Pt(18)
contact_p.add_run("Córdoba, Argentina  |  caropilardiaz@gmail.com  |  +54 351 155-428-714  |  linkedin.com/in/carolinadelpilardíaz").font.size = Pt(10)

date_p = doc.add_paragraph()
date_p.paragraph_format.space_after = Pt(14)
date_p.add_run("[Fecha]")

addr_p = doc.add_paragraph()
addr_p.paragraph_format.space_after = Pt(14)
addr_p.add_run("[Nombre del/de la responsable de contratación]\n[Empresa]\n[Ciudad, País]")

greeting_p = doc.add_paragraph()
greeting_p.paragraph_format.space_after = Pt(10)
greeting_p.add_run("Dear Hiring Team,")

body = [
    "I am writing to express my interest in the [Job Title] position at [Company]. As an Enterprise AI "
    "Solution Architect and Staff Statistician (Economist) with over 20 years of experience, I have spent "
    "the last several years steering corporate data strategies and engineering high-impact AI/ML ecosystems "
    "for global enterprises across Fintech, AgTech, and Healthcare — sectors where I believe my background "
    "aligns closely with the challenges your team is solving.",

    "In my current role as Staff AI Engineer / Architect at IT Patagonia, I designed and led a proprietary "
    "Multi-Agent Framework that now serves as the production blueprint for autonomous AI applications at "
    "top-tier financial institutions, including Banco Galicia and ICBC. My work centers on architecting "
    "provider-agnostic LLM systems — capable of dynamically switching between Gemini, GPT, and Claude — and "
    "on embedding rigorous statistical validation layers into agentic reasoning loops to prevent hallucinations "
    "in highly regulated environments. This combination of deep statistical rigor (Master's in Applied "
    "Statistics, background in Economics) and hands-on AI architecture is what allows me to bridge the gap "
    "between advanced model design and the compliance, governance, and explainability requirements that "
    "enterprise stakeholders demand.",

    "I am particularly drawn to [Company] because [personalizar: motivo específico — producto, misión, "
    "desafío técnico]. I would welcome the opportunity to bring my experience in enterprise AI architecture, "
    "AI governance, and technical leadership to your team, and to discuss how my track record of steering "
    "technical roadmaps — rather than only executing them — could contribute to your goals.",

    "Thank you for considering my application. I look forward to the possibility of speaking with you further.",
]

for para in body:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(para)

closing_p = doc.add_paragraph()
closing_p.paragraph_format.space_before = Pt(6)
closing_p.add_run("Sincerely,\nCarolina Del Pilar Díaz")

doc.save("Carolina_Diaz_Cover_Letter.docx")
print("Carta de presentación generada: Carolina_Diaz_Cover_Letter.docx")
