import docx
from docx.shared import Pt, Inches

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = docx.shared.RGBColor(0x33, 0x33, 0x33)

h = doc.add_paragraph()
r = h.add_run("Optimización de LinkedIn - Carolina Del Pilar Díaz")
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = docx.shared.RGBColor(0x00, 0x33, 0x66)

def add_section(doc, title, text_block):
    h_sec = doc.add_paragraph()
    h_sec.paragraph_format.space_before = Pt(14)
    r_s = h_sec.add_run(title)
    r_s.font.size = Pt(13)
    r_s.font.bold = True
    r_s.font.color.rgb = docx.shared.RGBColor(0x00, 0x33, 0x66)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text_block)

# 1. HEADLINE
add_section(
    doc,
    "1. Titular Principal (Headline)",
    "Staff AI Engineer & Solution Architect | Economist & Applied Statistician | Multi-Agent Systems "
    "(GenAI, RAG, MCP) | Python, R, SQL, Spark | AWS & Azure"
)

# 2. ABOUT
about_text = (
    "Staff AI Engineer and Solution Architect — Economist and Applied Statistician by training — with 20+ "
    "years spanning applied statistical research and hands-on production AI/ML engineering for global "
    "enterprises in Fintech, AgTech, and Healthcare.\n\n"
    "I spent over a decade leading statistical research and experimental design as Head of Statistical Research "
    "at INTA (Argentina's national agricultural research institute), then moved into enterprise AI — both "
    "architecting and personally building Multi-Agent Systems (GenAI, RAG, MCP) for regulated industries.\n\n"
    "As Staff AI Engineer / Architect at IT Patagonia, I designed a proprietary Multi-Agent Framework that now "
    "serves as the production blueprint for autonomous AI applications at Banco Galicia and ICBC — and I built "
    "and shipped it myself, end to end (LangGraph pipelines, RAG, FastAPI/React apps, AWS Lambda deployments). "
    "I build provider-agnostic LLM architectures that dynamically route between Gemini, GPT, and Claude, and I "
    "embed statistical validation and SHAP-based interpretability directly into agent reasoning loops to "
    "prevent hallucinations and meet compliance requirements (BCRA, CNV).\n\n"
    "What sets my approach apart is the combination of deep statistical rigor — a Master's in Applied Statistics "
    "and a background in Economics — with hands-on engineering of production-grade AI systems. I don't just "
    "design the architecture; I build it, and I embed the governance, validation, and observability layers that "
    "make GenAI systems auditable and trustworthy at enterprise scale.\n\n"
    "Databricks Certified Generative AI Engineer Associate, with a full Databricks Academy path completed in "
    "agentic AI: retrieval agents, single-agent applications, deployment & monitoring, evaluation & governance, "
    "and Unity Catalog.\n\n"
    "Core stack: Enterprise AI Architecture, Multi-Agent Systems (LangGraph, LangChain, LlamaIndex), RAG, MCP, "
    "AI Governance & XAI (SHAP), AWS (SageMaker, Bedrock, Fraud Detector), Azure AI, Databricks, Causal "
    "Inference, Econometrics, Bayesian Statistics.\n\n"
    "Languages: Spanish (Native), English (Professional Working Proficiency), Portuguese (Basic)."
)
add_section(doc, "2. Sección 'Acerca de' (About)", about_text)

# 3. LANGUAGES
lang_text = (
    "Configuración recomendada en el selector de LinkedIn:\n"
    "English → \"Professional Working Proficiency\".\n\n"
    "(Se descarta la etiqueta anterior \"B1/B2\" para que quede consistente con el CV y el Resume, "
    "que ya usan \"Professional Working Proficiency\".)"
)
add_section(doc, "3. Ajuste en la sección de Idiomas (Languages)", lang_text)

# 4. FEATURED
featured_text = (
    "Fijar arriba de todo (sección Featured) la insignia de:\n"
    "Databricks Certified Generative AI Engineer Associate — Issued 07/2026, valid through 07/2028 "
    "(Credential ID 186901264), con el link de verificación visible.\n\n"
    "Ya no es \"Candidate\": es la certificación aprobada, y debería ser el elemento más visible del perfil."
)
add_section(doc, "4. Sección Destacada (Featured)", featured_text)

doc.save("Word_LinkedIn_Carolina_Diaz.docx")
print("Archivo de LinkedIn generado con exito!")
